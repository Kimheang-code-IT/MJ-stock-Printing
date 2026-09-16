"""Generic table export — POST /api/v1/export/table (Excel + PDF downloads)."""

import io

import pytest

from tests.utils import admin_headers

_PAYLOAD = {
    "title": "Stock Report",
    "format": "xlsx",
    "subtitle": "Filtered rows",
    "columns": [
        {"key": "name", "label": "Name"},
        {"key": "qty", "label": "Qty"},
    ],
    "rows": [
        {"name": "Paracetamol", "qty": 3},
        {"name": "Glove", "qty": 10},
    ],
}


@pytest.mark.asyncio
async def test_export_table_xlsx_is_a_real_workbook(client):
    headers = await admin_headers(client)
    response = await client.post("/api/v1/export/table", json=_PAYLOAD, headers=headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith('.xlsx"')
    # XLSX is a zip container.
    assert response.content[:2] == b"PK"


@pytest.mark.asyncio
async def test_export_table_pdf_is_a_real_document(client):
    headers = await admin_headers(client)
    response = await client.post(
        "/api/v1/export/table", json={**_PAYLOAD, "format": "pdf"}, headers=headers
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].endswith('.pdf"')
    assert response.content[:4] == b"%PDF"


@pytest.mark.asyncio
async def test_export_table_xlsx_contains_rows_and_sum_formulas(client):
    """The spreadsheet must carry the actual rows plus real SUM totals."""
    from openpyxl import load_workbook

    headers = await admin_headers(client)
    payload = {
        "title": "របាយការណ៍លក់",
        "company": "Yoeun Sokhon Pharmacy",
        "format": "xlsx",
        "subtitle": "ចាប់ពី 2026-09-01 ដល់ 2026-09-17",
        "columns": [
            {"key": "no", "label": "ល.រ", "type": "number"},
            {"key": "product", "label": "ផលិតផល", "type": "text"},
            {"key": "amount", "label": "ចំនួនទឹកប្រាក់", "type": "money"},
        ],
        "rows": [
            {"no": 1, "product": "ស្រោមដៃ", "amount": 3},
            {"no": 2, "product": "Glove", "amount": 6},
        ],
    }
    response = await client.post("/api/v1/export/table", json=payload, headers=headers)
    assert response.status_code == 200, response.text

    sheet = load_workbook(io.BytesIO(response.content)).active
    values = list(sheet.iter_rows(values_only=True))
    cells = [cell for row in values for cell in row if cell is not None]
    # Real data, not just headers.
    assert "ស្រោមដៃ" in cells
    assert "Glove" in cells
    assert 3 in cells and 6 in cells
    # Totals are Excel formulas (so they match the report exactly).
    assert any(isinstance(cell, str) and cell.startswith("=SUM") for cell in cells)
    assert sheet.auto_filter.ref
    assert sheet.freeze_panes


@pytest.mark.asyncio
async def test_export_table_pdf_renders_khmer_without_error(client):
    """Khmer PDF must render a real multi-page-capable document (not blank)."""
    headers = await admin_headers(client)
    payload = {
        "title": "របាយការណ៍លក់",
        "company": "Yoeun Sokhon Pharmacy",
        "format": "pdf",
        "subtitle": "ចាប់ពី 2026-09-01 ដល់ 2026-09-17",
        "columns": [
            {"key": "date", "label": "កាលបរិច្ឆេទ", "type": "date"},
            {"key": "product", "label": "ផលិតផល", "type": "text"},
            {"key": "amount", "label": "ចំនួនទឹកប្រាក់", "type": "money"},
        ],
        "rows": [
            {"date": "2026-09-17", "product": "ស្រោមដៃ", "amount": 3},
            {"date": "2026-09-16", "product": "ប៉ារ៉ាសេតាម៉ុល ឈ្មោះវែង", "amount": 6},
        ],
    }
    response = await client.post("/api/v1/export/table", json=payload, headers=headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    assert response.content[:4] == b"%PDF"
    # A real report (embedded font + table) is far larger than a blank page.
    assert len(response.content) > 1000


@pytest.mark.asyncio
async def test_export_table_requires_authentication(client):
    response = await client.post("/api/v1/export/table", json=_PAYLOAD)
    assert response.status_code == 401


@pytest.mark.asyncio
async def test_export_table_rejects_unsupported_format(client):
    headers = await admin_headers(client)
    response = await client.post(
        "/api/v1/export/table", json={**_PAYLOAD, "format": "csv"}, headers=headers
    )
    assert response.status_code == 422
